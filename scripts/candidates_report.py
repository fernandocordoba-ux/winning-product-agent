"""Full candidates report: EVERY discovery candidate of a production run with every datum we have.

One row per product: discovery facts (90-day revenue when the lens returned it), filter status + reasons,
preliminary score, sustained growth, lens; deep analysis (WPS, confidence, WPS components, trend, creators,
videos, flags); decision (state, six dimensions, Amazon label, reasons, next actions); supplier estimates
(informative only). Missing values are N/A — nothing is estimated.

Writes reports/production/candidates/<run_id>.csv / .xlsx / .md (new files; the latest also as latest.*).
Uses the newest free rebuild (--report-only) of the run when there is one. Never a paid query.
"""
import csv
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(HERE))

import final_validation as FV  # noqa: E402
import discovery as D  # noqa: E402

NA = "N/A"
OUT = ROOT / "reports" / "production" / "candidates"
STATE_ES = {"READY_FOR_PRODUCT_VALIDATION": "Listo para validar", "PROMISING_NEEDS_VALIDATION": "Prometedor",
            "WATCHLIST": "En observación", "REJECT": "Descartado", "INSUFFICIENT_DATA": "Datos insuficientes"}


def _j(p):
    try:
        return json.loads(Path(p).read_text())
    except (OSError, TypeError, ValueError):
        return None


def _v(x):
    return NA if x is None or x == "" else x


def collect(root=ROOT):
    m, run_dir = FV.latest_production_run(root)
    if not m:
        return None, []
    ck = run_dir / "checkpoints"
    disc_ck = _j(ck / "discovery.json") or {}
    proc = _j(disc_ck.get("processed_file")) or {}
    # deep: newest rebuild checkpoint (rescore) when present
    deep = {}
    for mf in sorted(run_dir.parent.glob("*/manifest.json")):
        r = _j(mf) or {}
        if mf.parent == run_dir or r.get("report_only_source_run") == m.get("run_id"):
            for x in (_j(mf.parent / "checkpoints" / "deep_analysis.json") or {}).get("results") or []:
                deep[str(x.get("product_id"))] = x
    dec = _j(m["outputs"]["final_decision_json"]) or {}
    decs = {str(d["product_id"]): d for k in ("ready_products", "promising_products", "watchlist", "rejected_products",
                                              "insufficient_data") for d in dec.get(k) or []}
    sup_dir = root / "data" / "production" / "processed" / "suppliers"
    lens_by_raw = {}
    for q in disc_ck.get("queries") or []:
        env = _j(q.get("raw_file")) or {}
        lens_by_raw[str(q.get("raw_file"))] = env.get("lens")
    rows = []
    for p in proc.get("candidates") or []:
        f, c, o = p.get("facts") or {}, p.get("calculated") or {}, (p.get("source") or {}).get("original_record") or {}
        pid = str(f.get("product_id") or p.get("key"))
        d, x = deep.get(pid) or {}, decs.get(pid) or {}
        offers = []
        for of in sorted((sup_dir / pid).glob("offers_*.json")) if (sup_dir / pid).exists() else []:
            offers += (_j(of) or {}).get("offers") or []
        best = min((o2 for o2 in offers if o2.get("product_cost") is not None and o2.get("shipping_cost") is not None),
                   key=lambda o2: o2["product_cost"] + o2["shipping_cost"], default=None)
        amz = x.get("amazon_validation") or {}
        g = lambda k: D.parse_value(o.get(k), "num")[0]  # noqa: E731
        row = {
            "product_id": pid, "product_name": f.get("product_name"), "shop": f.get("shop_name"),
            "category": f.get("category") or d.get("category"), "lens": lens_by_raw.get(str((p.get("source") or {}).get("raw_file"))),
            "price_min": f.get("price_min"), "price_max": f.get("price_max"),
            "revenue_61_90d": g("gmv_prev2_30d"), "revenue_31_60d": f.get("gmv_prev_30d"), "revenue_last_30d": f.get("gmv_30d"),
            "growth_30d_pct": f.get("growth_30d"), "units_30d": f.get("units_30d"), "units_prev_30d": g("units_prev_30d"),
            "creators": f.get("creator_count"), "creator_growth_pct": f.get("creator_growth_pct"), "videos": f.get("video_count"),
            "shops_selling": f.get("shop_count"), "launch_date": f.get("launch_date"),
            "filter_status": c.get("filter_status"),
            "filter_reasons": "; ".join(str(r.get("rule")) for r in c.get("filter_reasons") or []),
            "preliminary_score": c.get("preliminary_wps"), "sustained_growth_90d": c.get("sustained_growth"),
            "deep_analyzed": bool(d), "deep_status": d.get("status"),
            "wps": d.get("wps"), "wps_confidence": d.get("confidence"),
            "trend": (d.get("trend_metrics") or {}).get("label") if d else None,
            "gmv_velocity_pct": ((d.get("trend_metrics") or {}).get("gmv") or {}).get("velocity_pct") if d else None,
            "video_sales_share_pct": (d.get("video_metrics") or {}).get("sales_share_pct") if d else None,
            "commission_pct": d.get("commission_pct"),
            "red_flags": "; ".join(sorted({str(r.get("code") or r.get("flag") or r) for r in d.get("red_flags") or []})) if d else None,
            "decision": STATE_ES.get(x.get("decision_state"), x.get("decision_state")),
            "decision_confidence": (x.get("decision_confidence") or {}).get("score"),
            "amazon": amz.get("label"), "amazon_reason": amz.get("reason"),
            "why_not": "; ".join(t.get("text", "") for t in x.get("why_it_did_not_pass") or []) if x else None,
            "next_actions": "; ".join(x.get("next_actions") or []) if x else None,
            "supplier_estimate": (f"{best.get('supplier_name')}: {best['product_cost']} + envío {best['shipping_cost']} USD"
                                  if best else None),
        }
        for k, comp in (d.get("wps_breakdown") or {}).items():
            row[f"wps_{k}"] = f"{comp.get('points')}/{comp.get('max')}"
        row["product_url"] = f.get("product_url") or (f"https://shop.tiktok.com/view/product/{pid}" if pid.isdigit() else None)
        rows.append(row)
    order = ["Listo para validar", "Prometedor", "En observación", "Datos insuficientes", "Descartado", None]
    rows.sort(key=lambda r: (not r["deep_analyzed"], order.index(r["decision"]) if r["decision"] in order else 9,
                             -(r["wps"] or 0), -(r["preliminary_score"] or 0)))
    return m, rows


def write(root=ROOT, out_dir=OUT):
    m, rows = collect(root)
    if not m:
        return None
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    cols = list(dict.fromkeys(k for r in rows for k in r))
    base = out_dir / m["run_id"]
    with open(base.with_suffix(".csv"), "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        w.writerows({k: _v(r.get(k)) for k in cols} for r in rows)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Candidatos"
        ws.append(cols)
        for r in rows:
            ws.append([_v(r.get(k)) if not isinstance(r.get(k), bool) else ("Sí" if r[k] else "No") for k in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F5C4A")
        ws.freeze_panes = "C2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(48, max(10, max(len(str(c.value or "")) for c in col[:40]) + 2))
        info = wb.create_sheet("Info")
        for k, v in (("Corrida", m["run_id"]), ("Configuración", m.get("config_version")),
                     ("Candidatos", len(rows)), ("Analizados a fondo", sum(r["deep_analyzed"] for r in rows)),
                     ("Nota", "N/A = dato no disponible (nunca estimado). Proveedores: solo referencia."),):
            info.append([k, v])
        wb.save(base.with_suffix(".xlsx"))
    except ImportError:
        pass
    L = [f"# Candidatos — corrida {m['run_id']}", "", f"Config {m.get('config_version')} · {len(rows)} candidatos · "
         f"{sum(r['deep_analyzed'] for r in rows)} analizados a fondo. Archivo completo: {base.name}.xlsx / .csv", "",
         "| # | Producto | Decisión | WPS | Preliminar | Ventas 90→60→30 | Unidades 30d | Amazon |", "|---|---|---|---|---|---|---|---|"]
    for i, r in enumerate(rows, 1):
        L.append(f"| {i} | {(r['product_name'] or '')[:60]} | {_v(r['decision'])} | {_v(r['wps'])} | {_v(r['preliminary_score'])} | "
                 f"{_v(r['revenue_61_90d'])} → {_v(r['revenue_31_60d'])} → {_v(r['revenue_last_30d'])} | {_v(r['units_30d'])} | {_v(r['amazon'])} |")
    base.with_suffix(".md").write_text("\n".join(L) + "\n")
    for ext in (".csv", ".xlsx", ".md"):
        if base.with_suffix(ext).exists():
            (out_dir / f"latest{ext}").write_bytes(base.with_suffix(ext).read_bytes())
    return {"run_id": m["run_id"], "rows": len(rows), "files": [str(base.with_suffix(e)) for e in (".csv", ".xlsx", ".md")]}


if __name__ == "__main__":
    print(write())
