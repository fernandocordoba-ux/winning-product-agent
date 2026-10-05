"""Calibration snapshot: frozen once per run, every discovered product, never overwritten. No network."""
import csv
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import candidates_report as CR  # noqa: E402


class Snapshot(unittest.TestCase):
    def test_written_once_with_low_scores_and_na(self):
        d = Path(tempfile.mkdtemp())
        m = {"run_id": "R1"}
        rows = [{"run_id": "R1", "product_id": "1", "wps": 73.09, "wps_floor": 73.09, "wps_coverage_pct": 100.0},
                {"run_id": "R1", "product_id": "2", "wps": None, "deep_analyzed": False}]     # never deep-analyzed
        p = CR.write_snapshot(m, rows, d)
        got = list(csv.DictReader(open(p)))
        self.assertEqual([r["product_id"] for r in got], ["1", "2"])
        self.assertEqual(got[1]["wps"], "N/A")                                 # missing stays N/A, never estimated
        self.assertEqual(set(got[0]), set(CR.SNAP_COLS))
        CR.write_snapshot(m, [{"run_id": "R1", "product_id": "9"}], d)         # frozen: a second call changes nothing
        self.assertEqual(len(list(csv.DictReader(open(p)))), 2)


if __name__ == "__main__":
    unittest.main()


class TwoViews(unittest.TestCase):
    def test_with_and_without_exclusions(self):
        from openpyxl import load_workbook
        d = Path(tempfile.mkdtemp())
        rows = [{"product_id": "1", "product_name": "Phone holder", "excluded": False, "wps": 60.0},
                {"product_id": "2", "product_name": "Power bank", "excluded": True,
                 "exclusion_reason": "exclusión operativa: power bank", "wps": None}]
        files = CR.write_two_views({"run_id": "R9"}, rows, d)
        self.assertEqual(len(files), 2)
        con = load_workbook(d / "R9-con-exclusiones.xlsx")["Candidatos"]
        sin = load_workbook(d / "R9-sin-exclusiones.xlsx")["Candidatos"]
        self.assertEqual(con.max_row, 2)                                   # header + the non-excluded product
        self.assertEqual(sin.max_row, 3)                                   # header + both
        names = [c.value for c in sin["A"]]
        self.assertIn("Power bank", names)
